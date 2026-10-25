import pymupdf
import pytest
from docx import Document

from services.document_parser import DocumentError, extract_text


def make_pdf(path, text):
    pdf = pymupdf.open()
    page = pdf.new_page()
    if text:
        page.insert_text((72, 72), text)
    pdf.save(str(path))
    pdf.close()


def make_docx(path, text):
    doc = Document()
    if text:
        doc.add_paragraph(text)
    doc.save(str(path))


def test_extract_pdf(tmp_path):
    f = tmp_path / "resume.pdf"
    make_pdf(f, "Jane Doe Python Developer")
    assert "Jane Doe" in extract_text(f)


def test_extract_docx(tmp_path):
    f = tmp_path / "resume.docx"
    make_docx(f, "John Smith Data Analyst")
    assert "John Smith" in extract_text(f)


def test_docx_table_text(tmp_path):
    f = tmp_path / "table.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "SQL"
    doc.save(str(f))
    assert "SQL" in extract_text(f)


def test_unsupported_type(tmp_path):
    f = tmp_path / "resume.txt"
    f.write_text("hello")
    with pytest.raises(DocumentError, match="Unsupported"):
        extract_text(f)


def test_missing_file(tmp_path):
    with pytest.raises(DocumentError, match="not found"):
        extract_text(tmp_path / "nope.pdf")


def test_zero_byte_file(tmp_path):
    f = tmp_path / "empty.pdf"
    f.write_bytes(b"")
    with pytest.raises(DocumentError, match="empty"):
        extract_text(f)


def test_blank_pdf(tmp_path):
    f = tmp_path / "blank.pdf"
    make_pdf(f, "")
    with pytest.raises(DocumentError, match="No readable text"):
        extract_text(f)


def test_blank_docx(tmp_path):
    f = tmp_path / "blank.docx"
    make_docx(f, "")
    with pytest.raises(DocumentError, match="No readable text"):
        extract_text(f)


def test_corrupted_pdf(tmp_path):
    f = tmp_path / "bad.pdf"
    f.write_bytes(b"this is not a real pdf")
    with pytest.raises(DocumentError, match="corrupted"):
        extract_text(f)


def test_corrupted_docx(tmp_path):
    f = tmp_path / "bad.docx"
    f.write_bytes(b"this is not a real docx")
    with pytest.raises(DocumentError, match="corrupted"):
        extract_text(f)