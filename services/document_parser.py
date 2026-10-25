"""Extract plain text from resume files (PDF and DOCX)."""
from pathlib import Path

import pymupdf
from docx import Document

MAX_FILE_SIZE_MB = 5
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class DocumentError(Exception):
    """Raised when a resume file cannot be read or is not valid."""


def _validate_file(path: Path) -> None:
    if not path.exists() or not path.is_file():
        raise DocumentError(f"File not found: {path.name}")
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise DocumentError("Unsupported file type. Please upload a PDF or DOCX file.")
    size_mb = path.stat().st_size / (1024 * 1024)
    if size_mb > MAX_FILE_SIZE_MB:
        raise DocumentError(f"File is too large (max {MAX_FILE_SIZE_MB} MB).")
    if path.stat().st_size == 0:
        raise DocumentError("The file is empty.")


def extract_text_from_pdf(path: Path) -> str:
    try:
        with pymupdf.open(path) as pdf:
            pages = [page.get_text() for page in pdf]
    except Exception as exc:
        raise DocumentError("Could not read the PDF. It may be corrupted.") from exc
    return "\n".join(pages)


def extract_text_from_docx(path: Path) -> str:
    try:
        doc = Document(str(path))
    except Exception as exc:
        raise DocumentError("Could not read the DOCX. It may be corrupted.") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_text(file_path) -> str:
    """Validate a resume file and return its cleaned text."""
    path = Path(file_path)
    _validate_file(path)
    if path.suffix.lower() == ".pdf":
        text = extract_text_from_pdf(path)
    else:
        text = extract_text_from_docx(path)
    text = text.strip()
    if not text:
        raise DocumentError(
            "No readable text found. The file may be empty or a scanned image."
        )
    return text